#!/usr/bin/env python3
"""Gera os .docx da documentacao a partir dos scripts gerar-*.ps1, sem Microsoft Word.

Interpreta o subconjunto de PowerShell usado nesses scripts (chamadas P/H1/H2/H3/Bul/Nota/
Exemplo/HR/Img/TableSimple e os comandos $sel.* da capa) e reproduz cada operacao do Word
(Selection.TypeText/TypeParagraph, fonte, paragrafo, bordas, tabelas, imagens, sumario)
com python-docx. O sumario e os numeros de pagina sao atualizados depois pelo LibreOffice
(ver atualizar_indices.py). Uso:

  python3 ferramentas/ps1_para_docx.py gerar-manual-uso.ps1 saida.docx [--img-dir PASTA]
"""
import re, sys, os, copy, argparse
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------- leitura do script
def split_statements(src):
    """Divide o script em instrucoes de topo, respeitando strings, () e {}."""
    stmts, buf, depth_p, depth_b = [], [], 0, 0
    i, n = 0, len(src)
    in_dq = in_sq = False
    in_comment = False
    while i < n:
        c = src[i]
        if in_comment:
            if c == '\n':
                in_comment = False
            else:
                i += 1; continue
        if in_dq:
            buf.append(c)
            if c == '`' and i + 1 < n:
                buf.append(src[i+1]); i += 2; continue
            if c == '"':
                if i + 1 < n and src[i+1] == '"':
                    buf.append('"'); i += 2; continue
                in_dq = False
            i += 1; continue
        if in_sq:
            buf.append(c)
            if c == "'":
                if i + 1 < n and src[i+1] == "'":
                    buf.append("'"); i += 2; continue
                in_sq = False
            i += 1; continue
        if c == '#' and (not buf or buf[-1] in ' \t\n;{}'):
            in_comment = True; i += 1; continue
        if c == '"': in_dq = True
        elif c == "'": in_sq = True
        elif c == '(': depth_p += 1
        elif c == ')': depth_p -= 1
        elif c == '{': depth_b += 1
        elif c == '}': depth_b -= 1
        if c in '\n;' and depth_p == 0 and depth_b == 0:
            s = ''.join(buf).strip()
            # "else" na linha seguinte continua o if
            if s:
                stmts.append(s)
            buf = []
        else:
            buf.append(c)
        i += 1
    s = ''.join(buf).strip()
    if s: stmts.append(s)
    # junta "if {...}" + "else {...}" separados por quebra de linha
    out = []
    for s in stmts:
        if s.startswith('else') and out and out[-1].startswith('if'):
            out[-1] += ' ' + s
        else:
            out.append(s)
    return out

class Tok:
    def __init__(self, s): self.s, self.i = s, 0
    def ws(self):
        while self.i < len(self.s) and self.s[self.i] in ' \t\r\n': self.i += 1
    def eof(self): self.ws(); return self.i >= len(self.s)
    def peek(self, k=1): return self.s[self.i:self.i+k]

ESC = {'0': '\0', 'a': '\a', 'b': '\b', 'e': '\x1b', 'f': '\f', 'n': '\n', 'r': '\r', 't': '\t', 'v': '\v'}

class Interp:
    def __init__(self, engine):
        self.e = engine
        self.vars = {'true': True, 'false': False, 'null': None}
        self.funcs = {}
        self.warn = []

    # ------------------------------------------------------------ expressoes
    def read_dq(self, t):
        assert t.s[t.i] == '"'; t.i += 1; out = []
        while True:
            c = t.s[t.i]
            if c == '`':
                nx = t.s[t.i+1]; out.append(nx); t.i += 2; continue  # crase usada como marcacao: mantem o caractere
            if c == '"':
                if t.s[t.i+1:t.i+2] == '"': out.append('"'); t.i += 2; continue
                t.i += 1; break
            if c == '$':
                m = re.match(r'\$([A-Za-z_][A-Za-z0-9_]*)', t.s[t.i:])
                if m:
                    name = m.group(1)
                    self.warn.append('interpolacao de $%s em string' % name)
                    v = self.vars.get(name.lower(), '')
                    out.append('' if v is None else str(v)); t.i += len(m.group(0)); continue
            out.append(c); t.i += 1
        return ''.join(out)

    def read_sq(self, t):
        t.i += 1; out = []
        while True:
            c = t.s[t.i]
            if c == "'":
                if t.s[t.i+1:t.i+2] == "'": out.append("'"); t.i += 2; continue
                t.i += 1; break
            out.append(c); t.i += 1
        return ''.join(out)

    def read_group(self, t, open_='(', close=')'):
        """retorna o texto entre parenteses balanceados (t.i aponta para o '(')"""
        depth, j, in_dq = 0, t.i, False
        while j < len(t.s):
            c = t.s[j]
            if in_dq:
                if c == '`': j += 2; continue
                if c == '"':
                    if t.s[j+1:j+2] == '"': j += 2; continue
                    in_dq = False
            elif c == '"': in_dq = True
            elif c == open_: depth += 1
            elif c == close:
                depth -= 1
                if depth == 0:
                    inner = t.s[t.i+1:j]; t.i = j + 1; return inner
            j += 1
        raise ValueError('parenteses sem fechamento: ' + t.s[t.i:t.i+80])

    def eval_array(self, inner):
        t = Tok(inner); items = []
        while not t.eof():
            items.append(self.eval_atom(t)); t.ws()
            if t.peek() == ',': t.i += 1
        return items

    def eval_atom(self, t):
        t.ws(); c = t.peek()
        if c == '"': return self.read_dq(t)
        if c == "'": return self.read_sq(t)
        if t.peek(2) == '@(': t.i += 1; return self.eval_array(self.read_group(t))
        if c == '(': return self.eval_expr(self.read_group(t))
        m = re.match(r'-?(0x[0-9A-Fa-f]+|\d+(\.\d+)?)', t.s[t.i:])
        if m:
            t.i += len(m.group(0)); v = m.group(0)
            return int(v, 16) if 'x' in v.lower() else (float(v) if '.' in v else int(v))
        m = re.match(r'\$([A-Za-z_][A-Za-z0-9_]*)((\.[A-Za-z]+(\([^()]*\))?)*)', t.s[t.i:])
        if m:
            t.i += len(m.group(0)); name, chain = m.group(1).lower(), m.group(2)
            if not chain: return self.vars.get(name)
            return self.eval_member(name, chain)
        m = re.match(r'[A-Za-z][A-Za-z0-9-]*', t.s[t.i:])
        if m:  # palavra solta (comando como argumento nao e usado; trata como string)
            t.i += len(m.group(0)); return m.group(0)
        raise ValueError('expressao nao reconhecida: ' + t.s[t.i:t.i+60])

    def eval_member(self, name, chain):
        m = re.match(r'\.(CentimetersToPoints|InchesToPoints)\((.*)\)$', chain)
        if name == 'word' and m:
            v = self.eval_expr(m.group(2))
            return v * 72 / 2.54 if m.group(1) == 'CentimetersToPoints' else v * 72
        m = re.match(r'\.Styles\.Item\((-?\d+)\)$', chain)
        if name == 'doc' and m: return ('style', int(m.group(1)))
        raise ValueError('membro nao suportado: $%s%s' % (name, chain))

    def eval_expr(self, s):
        s = s.strip()
        m = re.match(r'^if\s*\((.*?)\)\s*\{(.*?)\}\s*else\s*\{(.*?)\}$', s, re.S)
        if m:
            return self.eval_expr(m.group(2)) if self.eval_cond(m.group(1)) else self.eval_expr(m.group(3))
        if re.match(r'^[A-Za-z]', s) and not s.startswith('if'):
            # chamada de funcao como expressao, ex.: RGB 0xB9 0x1D 0x2E
            return self.call_line(s)
        t = Tok(s); v = self.eval_atom(t)
        if not t.eof():
            rest = s[t.i:].strip()
            m = re.match(r'^([+\-*/])\s*(.*)$', rest)
            if m:
                r = self.eval_expr(m.group(2)); op = m.group(1)
                return {'+': v + r if v is not None else r, '-': v - r, '*': v * r, '/': v / r}[op]
            raise ValueError('sobra na expressao: ' + rest)
        return v

    def eval_cond(self, s):
        m = re.match(r'^\s*(.+?)\s+-(eq|ne)\s+(.+?)\s*$', s)
        if m:
            a, b = self.eval_expr(m.group(1)), self.eval_expr(m.group(3))
            return (a == b) if m.group(2) == 'eq' else (a != b)
        return bool(self.eval_expr(s))

    # ------------------------------------------------------------ instrucoes
    def call_line(self, s):
        m = re.match(r'^([A-Za-z][A-Za-z0-9-]*)(.*)$', s, re.S)
        name, rest = m.group(1), m.group(2)
        args, t = [], Tok(rest)
        while not t.eof(): args.append(self.eval_atom(t))
        if name == 'RGB': return int(args[0] + args[1] * 256 + args[2] * 65536)
        if name in ('Write-Output', 'Write-Warning', 'Remove-Item'): return None
        if name in self.funcs: return self.run_func(name, args)
        if name in self.e.natives: return self.e.natives[name](self, args)
        raise ValueError('funcao desconhecida: ' + name)

    def run_func(self, name, args):
        params, body = self.funcs[name]
        saved = dict(self.vars)
        for k, (pname, default) in enumerate(params):
            self.vars[pname] = args[k] if k < len(args) else (self.eval_expr(default) if default is not None else None)
        for st in split_statements(body): self.exec(st)
        self.vars = saved

    def exec(self, s):
        s = s.strip().rstrip(';')
        if not s: return
        s = re.sub(r'\s*\|\s*Out-Null\s*$', '', s)
        if s.startswith('param('): return
        if 'New-Object' in s or re.match(r'^\$(word|doc|sel|sec|footer)\s*=', s): return
        m = re.match(r'^function\s+([A-Za-z0-9_]+)\s*\((.*?)\)\s*\{(.*)\}$', s, re.S)
        if m:
            name = m.group(1)
            params = []
            for p in [x.strip() for x in self.split_params(m.group(2)) if x.strip()]:
                pm = re.match(r'^\$([A-Za-z0-9_]+)\s*(=\s*(.*))?$', p, re.S)
                params.append((pm.group(1).lower(), pm.group(3)))
            if name in self.e.natives and name in ('Img', 'TableSimple', 'RGB'):
                return  # implementadas nativamente
            self.funcs[name] = (params, m.group(3))
            return
        m = re.match(r'^if\s*\((.*?)\)\s*\{(.*?)\}(\s*else\s*\{(.*)\})?$', s, re.S)
        if m:
            cond = m.group(1)
            if 'Test-Path' in cond: return
            branch = m.group(2) if self.eval_cond(cond) else (m.group(4) or '')
            for st in split_statements(branch): self.exec(st)
            return
        if s.startswith('for ') or s.startswith('for('):
            return self.e.raw_for(self, s)
        if s.startswith('$sel.') or s.startswith('$footer.') or s.startswith('$sec.') or s.startswith('$table') \
                or s.startswith('$doc.') or s.startswith('$word.') or s.startswith('$toc') or s.startswith('$tocRange') \
                or s.startswith('$tableRange') or s.startswith('[System') or s.startswith('$ErrorActionPreference') \
                or s.startswith('$range') or s.startswith('$nRows') or s.startswith('$nCols') or s.startswith('$shape') \
                or s.startswith('$cell') or s.startswith('$path') or s.startswith('$ratio'):
            return self.e.com(self, s)
        m = re.match(r'^\$([A-Za-z0-9_]+)\s*=\s*(.*)$', s, re.S)
        if m:
            self.vars[m.group(1).lower()] = self.eval_expr(m.group(2)); return
        return self.call_line(s)

    @staticmethod
    def split_params(s):
        out, buf, d = [], [], 0
        for c in s:
            if c in '(': d += 1
            if c in ')': d -= 1
            if c == ',' and d == 0: out.append(''.join(buf)); buf = []
            else: buf.append(c)
        out.append(''.join(buf)); return out

# ---------------------------------------------------------------- motor Word -> docx
def hexcolor(v):
    v = int(v); r, g, b = v & 255, (v >> 8) & 255, (v >> 16) & 255
    return '%02X%02X%02X' % (r, g, b)

NORMAL_FONT = dict(name=None, size=11, bold=False, italic=False, color=None)
NORMAL_PARA = dict(align=0, before=0, after=8, line=None, indent=0, bottom=None, left=None,
                   left_w=4, pbb=False, style=-1, bullet=False)

class Engine:
    def __init__(self, img_dir=None, default_font=None):
        self.doc = Document()
        st = self.doc.styles['Normal']
        if default_font:
            st.font.name = default_font
            st.element.rPr.rFonts.set(qn('w:eastAsia'), default_font)
        st.font.size = Pt(11)
        self.img_dir = img_dir
        self.font = dict(NORMAL_FONT)
        self.para = dict(NORMAL_PARA)
        self.runs = []
        self.pending_break = False
        self.table_rows, self.table_widths = None, {}
        self.footer_text = None
        # o documento novo do python-docx comeca com 0 paragrafos; ok
        self.natives = {'Img': self.img, 'TableSimple': self.table_simple}
        self.missing_imgs = []

    # --------- selecao
    def set_style(self, k):
        self.font = dict(NORMAL_FONT, name=self.font.get('name'))
        self.font['name'] = None
        self.para = dict(NORMAL_PARA, style=k)

    def type_text(self, text):
        self.runs.append((text, dict(self.font)))

    def type_paragraph(self):
        p = self.doc.add_paragraph()
        pr = self.para
        if pr['style'] == -2: p.style = self.doc.styles['Heading 1']
        elif pr['style'] == -3: p.style = self.doc.styles['Heading 2']
        if pr['bullet']:
            p.style = self.doc.styles['List Bullet']
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if pr['align'] == 1 else WD_ALIGN_PARAGRAPH.LEFT
        pf.space_before = Pt(pr['before']); pf.space_after = Pt(pr['after'])
        if pr['line']:
            pf.line_spacing = Pt(pr['line']); pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        if pr['bullet']:
            ind = pr['indent'] or Cm(0.6).pt
            pf.left_indent = Pt(ind + Cm(0.6).pt); pf.first_line_indent = Pt(-Cm(0.6).pt)
        else:
            pf.left_indent = Pt(pr['indent'])
        if pr['pbb'] or self.pending_break:
            pf.page_break_before = True; self.pending_break = False
        if pr['style'] in (-2, -3):
            pf.keep_with_next = True
        if pr['bottom'] is not None or pr['left'] is not None:
            pPr = p._p.get_or_add_pPr(); bdr = OxmlElement('w:pBdr')
            if pr['left'] is not None:
                e = OxmlElement('w:left'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(int(pr['left_w'])))
                e.set(qn('w:space'), '8'); e.set(qn('w:color'), hexcolor(pr['left'])); bdr.append(e)
            if pr['bottom'] is not None:
                e = OxmlElement('w:bottom'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4')
                e.set(qn('w:space'), '1'); e.set(qn('w:color'), hexcolor(pr['bottom'])); bdr.append(e)
            pPr.append(bdr)
        for text, f in self.runs:
            parts = text.split('\n')
            for k, part in enumerate(parts):
                r = p.add_run(part)
                if k < len(parts) - 1: r.add_break()
                self.apply_font(r, f)
        if not self.runs:
            # paragrafo vazio: a marca de paragrafo usa a fonte atual (afeta a altura da linha)
            rpr = p._p.get_or_add_pPr()
            r = p.add_run(''); self.apply_font(r, self.font)
        self.runs = []
        return p

    @staticmethod
    def apply_font(r, f):
        if f.get('name'):
            r.font.name = f['name']; r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), f['name'])
        r.font.size = Pt(f['size']); r.font.bold = bool(f['bold']); r.font.italic = bool(f['italic'])
        r.font.color.rgb = RGBColor.from_string(hexcolor(f['color'])) if f.get('color') is not None else RGBColor(0, 0, 0)

    # --------- comandos COM
    def com(self, it, s):
        m = re.match(r'^\$sel\.Style\s*=\s*(.*)$', s)
        if m:
            st = it.eval_expr(m.group(1)); self.flush_inline(); self.set_style(st[1]); return
        m = re.match(r'^\$sel\.Font\.(Name|Size|Bold|Italic|Color)\s*=\s*(.*)$', s)
        if m:
            k = m.group(1).lower(); v = it.eval_expr(m.group(2))
            self.font[{'name': 'name', 'size': 'size', 'bold': 'bold', 'italic': 'italic', 'color': 'color'}[k]] = v; return
        m = re.match(r'^\$sel\.ParagraphFormat\.(Alignment|SpaceAfter|SpaceBefore|LineSpacing|LeftIndent|PageBreakBefore)\s*=\s*(.*)$', s, re.S)
        if m:
            k = {'Alignment': 'align', 'SpaceAfter': 'after', 'SpaceBefore': 'before', 'LineSpacing': 'line',
                 'LeftIndent': 'indent', 'PageBreakBefore': 'pbb'}[m.group(1)]
            self.para[k] = it.eval_expr(m.group(2)); return
        m = re.match(r'^\$sel\.ParagraphFormat\.Borders\.Item\((\d)\)\.(LineStyle|Color|LineWidth)\s*=\s*(.*)$', s)
        if m:
            side = 'bottom' if m.group(1) == '3' else 'left'; v = it.eval_expr(m.group(3))
            key = '_' + side + '_color'
            if m.group(2) == 'LineStyle':
                if v == 0: self.para[side] = None
                else: self.para[side] = self.para.get(key, 0)
            elif m.group(2) == 'Color':
                self.para[key] = v
                if self.para.get(side) is not None or True: self.para[side] = v if self.para.get(side) is not None else self.para.get(side)
                # Word: cor so aparece se a borda estiver ligada; liga-se antes ou depois
                if self.para.get(side) is None and self.para.get('_' + side + '_on'): self.para[side] = v
            else:
                self.para['left_w'] = v
            if m.group(2) == 'LineStyle':
                self.para['_' + side + '_on'] = (v != 0)
                if v != 0: self.para[side] = self.para.get(key, 0)
            return
        if re.match(r'^\$sel\.TypeText\(', s):
            t = Tok(s[len('$sel.TypeText'):]); self.type_text(it.eval_expr(it.read_group(t))); return
        if s == '$sel.TypeParagraph()': self.type_paragraph(); return
        if s == '$sel.InsertBreak(7)':
            if self.runs: self.type_paragraph()
            self.pending_break = True; return
        if s == '$sel.Range.ListFormat.ApplyBulletDefault()': self.para['bullet'] = True; return
        if s == '$sel.Range.ListFormat.RemoveNumbers()': self.para['bullet'] = False; return
        if s.startswith('$toc = $doc.TablesOfContents.Add'): self.add_toc(); return
        m = re.match(r'^\$footer\.Range\.Text\s*=\s*(.*)$', s)
        if m: self.footer_text = it.eval_expr(m.group(1)); return
        if s.startswith('$footer.Range.Fields.Add'): self.add_footer(); return
        m = re.match(r'^\$sec\.PageSetup\.Page(Width|Height)\s*=\s*(.*)$', s)
        if m:
            sec = self.doc.sections[0]; v = Pt(it.eval_expr(m.group(2)))
            if m.group(1) == 'Width': sec.page_width = v
            else: sec.page_height = v
            return
        if s.startswith('$table = $doc.Tables.Add('):
            self.table_rows = it.vars.get('rows'); self.table_widths = {}; return
        m = re.match(r'^\$table\.Columns\.Item\((\d+)\)\.Width\s*=\s*(.*)$', s)
        if m: self.table_widths[int(m.group(1))] = it.eval_expr(m.group(2)) / 72 * 2.54; return
        if s.startswith('$sel.EndKey(6)'):
            if self.table_rows is not None:
                n = len(self.table_rows[0])
                self.make_table(self.table_rows, [self.table_widths.get(i + 1, 4) for i in range(n)], font=None)
                self.table_rows = None
            return
        # o resto (criar Word, salvar, fechar, bordas da tabela, etc.) nao produz conteudo
        return

    def raw_for(self, it, s):
        return  # lacos de preenchimento de tabela: tratados em make_table

    def flush_inline(self):
        pass

    # --------- blocos especiais
    def add_toc(self):
        p = self.doc.add_paragraph()
        r = p.add_run()
        b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); r._r.append(b)
        r2 = p.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve')
        it.text = ' TOC \\o "1-2" \\h \\z \\u '; r2._r.append(it)
        r3 = p.add_run(); sp = OxmlElement('w:fldChar'); sp.set(qn('w:fldCharType'), 'separate'); r3._r.append(sp)
        p.add_run('Sumario - atualize o campo para gerar.')
        r4 = p.add_run(); e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r4._r.append(e)

    def add_footer(self):
        ft = self.doc.sections[0].footer
        p = ft.paragraphs[0] if ft.paragraphs else ft.add_paragraph()
        r = p.add_run(self.footer_text or ''); r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string('6A6A70')
        r = p.add_run()
        for typ, txt in (('begin', None), (None, ' PAGE '), ('separate', None), (None, '1'), ('end', None)):
            if typ:
                el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), typ); r._r.append(el)
            elif txt == ' PAGE ':
                el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = txt; r._r.append(el)
            else:
                t = OxmlElement('w:t'); t.text = txt; r._r.append(t)
        r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string('6A6A70')

    def img(self, it, args):
        filename, caption = args[0], args[1] if len(args) > 1 else None
        width = args[2] if len(args) > 2 else 5.6
        path = os.path.join(self.img_dir or '.', filename)
        if not os.path.exists(path):
            self.missing_imgs.append(filename); return
        if self.runs: self.type_paragraph()
        p = self.doc.add_paragraph(); pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER; pf.space_before = Pt(6); pf.space_after = Pt(2)
        if self.pending_break: pf.page_break_before = True; self.pending_break = False
        p.add_run().add_picture(path, width=Inches(width))
        # contorno cinza de 0,75 pt, como no script original
        spPr = p._p.xpath('.//pic:spPr')[0]
        ln = OxmlElement('a:ln'); ln.set('w', str(int(0.75 * 12700)))
        sf = OxmlElement('a:solidFill'); c = OxmlElement('a:srgbClr'); c.set('val', 'D1D5DB'); sf.append(c); ln.append(sf)
        spPr.append(ln)
        if caption:
            p = self.doc.add_paragraph(); pf = p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER; pf.space_after = Pt(14); pf.space_before = Pt(0)
            r = p.add_run(caption); f = dict(name='Montserrat', size=9.5, bold=False, italic=True, color=it.vars.get('colmuted'))
            self.apply_font(r, f)
        self.para['align'] = 0
        self.para['before'] = 6; self.para['after'] = 14 if caption else 2
        self.font['italic'] = False

    def table_simple(self, it, args):
        self.make_table(args[0], args[1], font='Montserrat')
        self.para.update(dict(NORMAL_PARA)); self.para['after'] = 10
        self.type_paragraph()

    def make_table(self, rows, widths_cm, font):
        if self.runs: self.type_paragraph()
        nR, nC = len(rows), len(rows[0])
        t = self.doc.add_table(rows=nR, cols=nC)
        tblPr = t._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            e = OxmlElement('w:' + side); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), 'D1D5DB'); borders.append(e)
        tblPr.append(borders)
        t.autofit = False
        for c in range(nC):
            for cell in t.columns[c].cells: cell.width = Cm(widths_cm[c])
        for r in range(nR):
            for c in range(nC):
                cell = t.cell(r, c); cell.text = ''
                p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
                run = p.add_run(str(rows[r][c]))
                self.apply_font(run, dict(name=font, size=9.5, bold=(r == 0), italic=False, color=0xFFFFFF if r == 0 else 0x1A1A1A))
                if r == 0:
                    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
                    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), '1A1A1A'); tcPr.append(shd)
            if r == 0:
                trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); trPr.append(h)
        return t

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('ps1'); ap.add_argument('out'); ap.add_argument('--img-dir')
    a = ap.parse_args()
    src = open(a.ps1, encoding='utf-8-sig').read()
    eng = Engine(img_dir=a.img_dir, default_font='Montserrat' if 'Montserrat' in src else None)
    it = Interp(eng)
    for st in split_statements(src):
        try:
            it.exec(st)
        except Exception as ex:
            raise SystemExit('ERRO em: %s\n  -> %s' % (st[:200], ex))
    if eng.runs: eng.type_paragraph()
    eng.doc.save(a.out)
    for w in sorted(set(it.warn)): print('AVISO:', w)
    if eng.missing_imgs: print('IMAGENS AUSENTES:', len(eng.missing_imgs), eng.missing_imgs[:5])
    print('SALVO:', a.out)

if __name__ == '__main__':
    main()
